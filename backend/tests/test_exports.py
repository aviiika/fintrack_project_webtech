"""CSV and Excel export endpoints."""
import io


def test_csv_export_has_header_and_rows(seeded_client):
    client, auth_headers = seeded_client
    r = client.get("/api/v1/export/transactions.csv", headers=auth_headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"]

    text = r.content.decode("utf-8")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    assert lines[0].startswith("Date,Description,Category")
    assert len(lines) > 1  # header + at least one transaction
    assert "Zomato" in text


def test_csv_export_respects_month_filter(seeded_client):
    client, auth_headers = seeded_client
    # Seeded data is all May 2025, so a different month must come back empty.
    r = client.get("/api/v1/export/transactions.csv?month=2030-01", headers=auth_headers)
    assert r.status_code == 200
    lines = [ln for ln in r.content.decode().splitlines() if ln.strip()]
    assert len(lines) == 1  # header only


def test_xlsx_export_is_a_valid_workbook(seeded_client):
    client, auth_headers = seeded_client
    r = client.get("/api/v1/export/report.xlsx", headers=auth_headers)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Transactions", "By Category", "Monthly Summary"]
    assert wb["Transactions"].max_row > 1
    assert wb["By Category"].max_row > 1


def test_export_requires_auth(client):
    assert client.get("/api/v1/export/transactions.csv").status_code == 401
    assert client.get("/api/v1/export/report.xlsx").status_code == 401
